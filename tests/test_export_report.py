from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from roi_model.analysis import analyze
from roi_model.assumptions import load_case
from roi_model.export_tabular import write_csv, write_xlsx
from roi_model.report import render_html, render_markdown, render_one_pager

ROOT = Path(__file__).parents[1]


def test_csv_xlsx_and_report_labels(tmp_path: Path) -> None:
    analysis = analyze(load_case(ROOT / "assumptions" / "refund_desk.yaml"), draws=25, seed=2)
    write_csv(analysis, tmp_path)
    csv_text = (tmp_path / "results.csv").read_text(encoding="utf-8")
    inputs_csv = (tmp_path / "inputs.csv").read_text(encoding="utf-8")
    assert "quality_adjusted_savings_per_item" in csv_text
    assert "illustrative" in inputs_csv
    assert "[EST]" in inputs_csv
    write_xlsx(analysis, tmp_path / "model.xlsx")
    book = load_workbook(tmp_path / "model.xlsx")
    labels = {book["Model"].cell(row, 1).value: book["Model"].cell(row, 2).value for row in range(2, 14)}
    formulas = list(labels.values())
    assert book["Inputs"].cell(1, 4).value == "illustrative"
    assert book["Inputs"].cell(2, 4).value == "illustrative"
    assert "[EST]" in str(book["Inputs"].cell(2, 5).value)
    assert labels["quality_adjusted_savings"] == "=B3+B4-(B2+B5)-B6"
    assert labels["net_monthly_steady"] == "=B8*B7-MonthlyRunCost"
    assert all(isinstance(value, str) and value.startswith("=") for value in formulas)
    joined = "\n".join(str(value) for value in formulas)
    assert "ManualMinutes" in joined
    assert "never" in joined
    markdown = render_markdown(
        analysis,
        generated_at="2026-10-17T00:00:00+00:00",
        command="python -m roi_model.cli run",
        git_revision="abc",
    )
    assert "[EST]" in markdown
    assert "simulated" in markdown.lower()
    assert "Measured-cost slot" in markdown
    assert "illustrative" in markdown
    assert markdown.startswith("Illustrative numbers, replace with your own")
    html = render_html(analysis, markdown, "<svg></svg>")
    assert "<svg" in html
    assert "Puerto Gris" in html
    english = render_one_pager(
        analysis,
        lang="en",
        generated_at="2026-10-17T00:00:00+00:00",
        command="cmd",
        git_revision="abc",
    )
    spanish = render_one_pager(
        analysis,
        lang="es",
        generated_at="2026-10-17T00:00:00+00:00",
        command="cmd",
        git_revision="abc",
    )
    assert english.startswith("Illustrative numbers, replace with your own")
    assert spanish.startswith("Illustrative numbers, replace with your own")
    assert "Simulated" in english
    assert "illustrative [EST]" in english
    assert "Mirá" in spanish
    assert "[EST]" in spanish
    weekly = analyze(load_case(ROOT / "assumptions" / "weekly_report.yaml"), draws=15, seed=3)
    write_xlsx(weekly, tmp_path / "weekly.xlsx")
    assert load_workbook(tmp_path / "weekly.xlsx")["Model"]["B2"].value.startswith("=")
